"""Local Windows Teams caption OCR. Python 3.11+; see README.md.
No Teams integration, audio capture, network calls, or mouse tracking after selection.
"""
from __future__ import annotations

import argparse
import csv
import ctypes
import difflib
import io
import json
import os
from pathlib import Path
import queue
import re
import shutil
import statistics
import subprocess
import sys
import threading
import time
from dataclasses import dataclass, field
from datetime import datetime


def normalize(text: str) -> str:
    return " ".join(re.findall(r"[^\W_]+", text.casefold(), flags=re.UNICODE))


def similarity(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, normalize(a), normalize(b), autojunk=False).ratio()


def overlap_count(old: str, new: str) -> int:
    """Conservative suffix/prefix match: never suppress a lone repeated yes/no."""
    a, b = old.split(), new.split()
    for count in range(min(len(a), len(b), 100), 2, -1):
        x, y = normalize(" ".join(a[-count:])), normalize(" ".join(b[:count]))
        if x == y or (count >= 5 and similarity(x, y) >= .94):
            return count
    return 0


def merge_revision(old: str, new: str) -> str:
    """Update an existing card, preserving text that rolled out of view."""
    a, b = normalize(old), normalize(new)
    if a == b:
        return old  # do not oscillate on punctuation-only OCR differences
    if b.startswith(a):
        return new
    if a.startswith(b) or (len(b) >= 12 and b in a):
        return old  # a clipped/briefly incomplete capture is not a deletion
    if len(old.split()) == len(new.split()) and similarity(old, new) >= .90:
        return new
    overlap = overlap_count(old, new)
    if overlap:
        return old + (" " + " ".join(new.split()[overlap:]) if len(new.split()) > overlap else "")
    if similarity(old, new) >= .70:
        return new  # revise in place; do not append OCR corrections as speech
    return old


@dataclass
class Row:
    text: str
    y: float
    height: float
    words: list[dict] = field(default_factory=list)


@dataclass
class Card:
    speaker: str
    text: str
    y: float = 0
    key: int | None = None


class CaptionParser:
    """Teams black-panel layout: smaller header, then one or more larger lines.

    Select the TEXT COLUMN, excluding circular avatars. Optional --trim-left
    handles a region that includes avatars. Unknown top-of-panel continuation
    lines are skipped until a speaker header is visible.
    """
    def __init__(self, speakers=()):
        self.names = list(speakers)

    def name_for(self, row: Row, body_height: float) -> str | None:
        # Drop status icons separated from the name by a larger horizontal gap.
        words = row.words
        kept = []
        for word in words:
            if kept and word['x'] - (kept[-1]['x'] + kept[-1]['w']) > row.height * .65:
                break
            kept.append(word)
        text = " ".join(w['text'] for w in kept) if kept else row.text
        text = re.sub(r"[^\w, .’'\-]+$", "", text).strip()
        text = re.sub(r"\s+\([^)]+\)$", "", text).strip()
        if not text:
            return None
        for known in self.names:
            if normalize(text) == normalize(known):
                return known
            if len(normalize(text)) >= 8 and similarity(text, known) >= .88 and row.height < body_height:
                return known
        parts = text.replace(',', ' ').split()
        if not 2 <= len(parts) <= 6 or len(text) > 65:
            return None
        # Proper-name grammar plus font size prevents caption sentences with
        # commas from being treated as people. Explicit --speaker overrides help
        # with all-caps, lowercase, or unusual Teams display names.
        if not all(re.fullmatch(r"[^\W\d_]+(?:[’'.-][^\W\d_]+)*\.?", p) for p in parts):
            return None
        if not all(p[0].isupper() or p.lower() in {'de', 'van', 'von', 'da', 'del'} for p in parts):
            return None
        if row.height >= body_height * .96:
            return None
        self.names.append(text)
        return text

    def parse(self, rows: list[Row]) -> list[Card]:
        if not rows:
            return []
        # Caption body occupies most lines; upper half avoids small header bias.
        heights = sorted(r.height for r in rows if r.height > 0)
        body_height = statistics.median(heights[len(heights)//2:]) if heights else 1
        cards = []
        current = None
        for row in rows:
            name = self.name_for(row, body_height)
            if name:
                current = Card(name, '', row.y)
                cards.append(current)
            elif current:
                current.text = (current.text + ' ' + row.text).strip()
        return [c for c in cards if c.text]


@dataclass
class Segment:
    key: int
    text: str


@dataclass
class Block:
    speaker: str
    start: str
    end: str
    segments: list[Segment] = field(default_factory=list)

    @property
    def text(self):
        result = ''
        for segment in self.segments:
            # Adjacent *distinct* cards can legitimately repeat phrases. Only
            # strip a substantial boundary overlap, not a whole repeated card.
            n = overlap_count(result, segment.text) if result else 0
            if n == len(segment.text.split()):
                n = 0
            addition = ' '.join(segment.text.split()[n:])
            result = (result + ' ' + addition).strip()
        return result


class Transcript:
    def __init__(self, gap_seconds=8.0):
        self.blocks: list[Block] = []
        self.previous: list[Card] = []
        self.locations: dict[int, tuple[Block, Segment]] = {}
        self.next_key = 1
        self.last_change = None
        self.blank_since = None
        self.gap_seconds = gap_seconds

    @staticmethod
    def score(a: Card, b: Card) -> float:
        if a.speaker != b.speaker:
            return 0
        x, y = normalize(a.text), normalize(b.text)
        if x == y:
            return 2
        if min(len(x), len(y)) >= 3 and (x.startswith(y + ' ') or y.startswith(x + ' ')):
            return 1.7
        if min(len(x), len(y)) >= 8 and (x in y or y in x):
            return 1.6
        ratio = similarity(x, y)
        if ratio >= .72 and min(len(x), len(y)) >= 8:
            return ratio
        n = overlap_count(a.text, b.text)
        if n and n >= min(len(a.text.split()), len(b.text.split())) * .4:
            return .75
        return 0

    def consume(self, cards: list[Card], now: datetime) -> bool:
        if not cards:
            if self.blank_since is None:
                self.blank_since = now
            elif (now - self.blank_since).total_seconds() > self.gap_seconds:
                self.previous = []
            return False
        self.blank_since = None
        old = self.previous
        # Weighted order-preserving alignment handles scrolling cards and
        # duplicate names. Each prior card may match at most one current card.
        m, n = len(old), len(cards)
        dp = [[0.0] * (n + 1) for _ in range(m + 1)]
        steps = {}
        for i in range(1, m + 1):
            for j in range(1, n + 1):
                score = self.score(old[i-1], cards[j-1])
                choices = [(dp[i-1][j], 'old'), (dp[i][j-1], 'new')]
                if score:
                    choices.append((dp[i-1][j-1] + score, 'match'))
                dp[i][j], steps[i, j] = max(choices, key=lambda v: v[0])
        matched = {}
        i, j = m, n
        while i and j:
            step = steps[i, j]
            if step == 'match':
                matched[j-1] = old[i-1].key
                i -= 1
                j -= 1
            elif step == 'old':
                i -= 1
            else:
                j -= 1
        changed = False
        stamp = now.astimezone().isoformat(timespec='seconds')
        last_matched = max(matched, default=-1)
        opened_new = False
        for idx, card in enumerate(cards):
            if idx in matched:
                card.key = matched[idx]
                block, segment = self.locations[card.key]
                updated = merge_revision(segment.text, card.text)
                if updated != segment.text:
                    segment.text = updated
                    block.end = stamp
                    changed = True
            elif idx < last_matched:
                # A previously visible card briefly missed by OCR may return
                # above a matched card. Do not insert stale speech out of order.
                continue
            else:
                card.key = self.next_key
                self.next_key += 1
                idle = self.last_change is not None and (now - self.last_change).total_seconds() >= self.gap_seconds
                if not self.blocks or self.blocks[-1].speaker != card.speaker or (idle and not opened_new):
                    self.blocks.append(Block(card.speaker, stamp, stamp))
                block = self.blocks[-1]
                block.end = stamp
                segment = Segment(card.key, card.text)
                block.segments.append(segment)
                opened_new = True
                self.locations[card.key] = (block, segment)
                changed = True
        self.previous = [c for c in cards if c.key is not None]
        if changed:
            self.last_change = now
        return changed

    def render(self) -> str:
        lines = ['Teams caption transcript',
                 'Timestamps are local screen-observation times, not audio timecodes.', '']
        for block in self.blocks:
            lines.extend([f'[{block.start}] {block.speaker}', block.text, ''])
        return '\n'.join(lines)


def atomic_text(path: Path, text: str):
    temp = path.with_suffix(path.suffix + '.tmp')
    with temp.open('w', encoding='utf-8', newline='\n') as handle:
        handle.write(text)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temp, path)


class Output:
    def __init__(self, folder: Path):
        folder.mkdir(parents=True, exist_ok=True)
        self.base = folder / ('Teams_' + datetime.now().strftime('%Y-%m-%d_%H-%M-%S_%f'))
        self.journal = self.base.with_suffix('.jsonl')
        self.docx_warning = ''

    def save(self, transcript: Transcript):
        # Clean TXT is atomically refreshed to permit correction of an evolving
        # OCR card; no duplicate revisions are appended to the clean transcript.
        atomic_text(self.base.with_suffix('.txt'), transcript.render())
        # Append-only recovery journal holds the latest state of changed blocks.
        states = getattr(self, '_states', {})
        with self.journal.open('a', encoding='utf-8') as handle:
            for idx, block in enumerate(transcript.blocks):
                state = {'block': idx, 'speaker': block.speaker, 'start': block.start,
                         'end': block.end, 'text': block.text}
                if states.get(idx) != state:
                    handle.write(json.dumps(state, ensure_ascii=False) + '\n')
                    states[idx] = state
            handle.flush()
            os.fsync(handle.fileno())
        self._states = states

    def docx(self, transcript: Transcript):
        try:
            from docx import Document
            from docx.shared import Pt
            doc = Document()
            doc.styles['Normal'].font.name = 'Calibri'
            doc.styles['Normal'].font.size = Pt(11)
            doc.add_heading('Teams caption transcript', 0)
            doc.add_paragraph('Local screen-observation timestamps; OCR text may contain errors.')
            for block in transcript.blocks:
                paragraph = doc.add_paragraph()
                paragraph.add_run(f'[{block.start}] {block.speaker}').bold = True
                doc.add_paragraph(block.text)
            temp = self.base.with_suffix('.tmp.docx')
            doc.save(temp)
            os.replace(temp, self.base.with_suffix('.docx'))
            self.docx_warning = ''
        except (ImportError, OSError) as exc:
            self.docx_warning = f'DOCX unavailable: {exc}. TXT and JSONL remain saved.'


def find_tesseract(explicit=None) -> str:
    candidates = [explicit, os.getenv('TESSERACT_CMD'), shutil.which('tesseract'),
                  r'C:\Program Files\Tesseract-OCR\tesseract.exe',
                  str(Path(os.getenv('LOCALAPPDATA', '')) / 'Programs/Tesseract-OCR/tesseract.exe')]
    for path in candidates:
        if path and Path(path).is_file():
            return str(path)
    raise RuntimeError('Tesseract not found. Install it, or use --tesseract "C:\\path\\tesseract.exe".')


def ocr_rows(image, executable: str, language='eng', trim_left=0) -> list[Row]:
    from PIL import ImageOps
    if trim_left:
        if trim_left >= image.width - 40:
            raise ValueError('--trim-left leaves too little caption width.')
        image = image.crop((trim_left, 0, image.width, image.height))
    gray = ImageOps.grayscale(image)
    # Teams dark captions become dark-on-light for Tesseract.
    if statistics.mean(gray.resize((16, 16)).getdata()) < 128:
        gray = ImageOps.invert(gray)
    scale = 3
    gray = ImageOps.autocontrast(gray).resize((gray.width * scale, gray.height * scale))
    png = io.BytesIO()
    gray.save(png, format='PNG')
    result = subprocess.run(
        [executable, 'stdin', 'stdout', '-l', language, '--psm', '6', 'tsv'],
        input=png.getvalue(), capture_output=True, timeout=8,
        creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == 'win32' else 0)
    if result.returncode:
        raise RuntimeError(result.stderr.decode('utf-8', errors='replace').strip())
    groups = {}
    for record in csv.DictReader(io.StringIO(result.stdout.decode('utf-8', errors='replace')), delimiter='\t', quoting=csv.QUOTE_NONE):
        text = (record.get('text') or '').strip()
        if record['level'] != '5' or not text or float(record['conf']) < 20:
            continue
        x, y, w, h = (int(record[k]) for k in ('left', 'top', 'width', 'height'))
        if y <= 2 or y + h >= gray.height - 2:
            continue  # incomplete glyphs at selected region boundaries
        key = tuple(record[k] for k in ('block_num', 'par_num', 'line_num'))
        groups.setdefault(key, []).append({'text': text, 'x': x/scale, 'y': y/scale, 'w': w/scale, 'h': h/scale})
    rows = []
    for words in groups.values():
        words.sort(key=lambda word: word['x'])
        letters = [w for w in words if any(c.isalpha() for c in w['text'])]
        if not letters:
            continue
        height = statistics.median(w['h'] for w in letters)
        rows.append(Row(' '.join(w['text'] for w in words), min(w['y'] for w in words), height, words))
    return sorted(rows, key=lambda row: row.y)


def dpi_awareness():
    if sys.platform == 'win32':
        try:
            ctypes.windll.user32.SetProcessDpiAwarenessContext(ctypes.c_void_p(-4))
        except (AttributeError, OSError):
            try:
                ctypes.windll.shcore.SetProcessDpiAwareness(2)
            except (AttributeError, OSError):
                pass


def select_region(root):
    import tkinter as tk
    from PIL import Image, ImageTk
    import mss
    root.withdraw()
    root.update()
    time.sleep(.25)
    with mss.mss() as capture:
        bounds = dict(capture.monitors[0])
        shot = capture.grab(bounds)
        frozen = Image.frombytes('RGB', shot.size, shot.rgb)
    window = tk.Toplevel(root)
    window.overrideredirect(True)
    window.attributes('-topmost', True)
    window.geometry(f"{bounds['width']}x{bounds['height']}+0+0")
    window.update_idletasks()
    if sys.platform == 'win32':
        from ctypes import wintypes
        api = ctypes.windll.user32
        api.GetParent.argtypes = [wintypes.HWND]
        api.GetParent.restype = wintypes.HWND
        api.SetWindowPos.argtypes = [wintypes.HWND, wintypes.HWND, ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int, wintypes.UINT]
        hwnd = api.GetParent(window.winfo_id()) or window.winfo_id()
        api.SetWindowPos(hwnd, wintypes.HWND(-1), bounds['left'], bounds['top'], bounds['width'], bounds['height'], 0x0040)
    canvas = tk.Canvas(window, width=bounds['width'], height=bounds['height'], highlightthickness=0, cursor='crosshair')
    canvas.pack(fill='both', expand=True)
    photo = ImageTk.PhotoImage(frozen, master=window)
    canvas.create_image(0, 0, image=photo, anchor='nw')
    # Help is shown in the control window before capture; no banner covers captions.
    result, origin, rectangle = [], [], []
    def down(event):
        origin[:] = [event.x, event.y]
        if rectangle:
            canvas.delete(rectangle.pop())
        rectangle.append(canvas.create_rectangle(event.x, event.y, event.x, event.y, outline='#00dcff', width=3))
    def move(event):
        if origin:
            canvas.coords(rectangle[0], *origin, event.x, event.y)
    def up(event):
        if not origin:
            return
        left, right = sorted((origin[0], event.x))
        top, bottom = sorted((origin[1], event.y))
        if right-left < 120 or bottom-top < 40:
            return
        result.append({'left': bounds['left']+left, 'top': bounds['top']+top,
                       'width': right-left, 'height': bottom-top})
        window.destroy()
    canvas.bind('<ButtonPress-1>', down)
    canvas.bind('<B1-Motion>', move)
    canvas.bind('<ButtonRelease-1>', up)
    window.bind('<Escape>', lambda _: window.destroy())
    window.focus_force()
    window.grab_set()
    root.wait_window(window)
    root.deiconify()
    return result[0] if result else None


class App:
    def __init__(self, root, args, executable):
        import tkinter as tk
        from tkinter import ttk
        from tkinter.scrolledtext import ScrolledText
        self.root, self.args, self.executable = root, args, executable
        self.events = queue.Queue()
        self.stop_event = threading.Event()
        self.thread = None
        self.closing = False
        self.output = None
        root.title('Teams Caption Capture — Local OCR')
        root.geometry('740x520')
        ttk.Label(root, text='1. Show Teams captions.  2. Select the text column, excluding the initials circles.\n'
                  'Include names and complete caption lines. Esc cancels selection.\n'
                  'After selection, keep this window outside the caption region.', padding=12).pack(anchor='w')
        bar = ttk.Frame(root, padding=8)
        bar.pack(fill='x')
        self.start_button = ttk.Button(bar, text='Select region & start', command=self.start)
        self.start_button.pack(side='left', padx=4)
        self.stop_button = ttk.Button(bar, text='Stop & export', state='disabled', command=self.stop)
        self.stop_button.pack(side='left', padx=4)
        self.folder_button = ttk.Button(bar, text='Open output folder', state='disabled', command=self.open_folder)
        self.folder_button.pack(side='left', padx=4)
        self.status = tk.StringVar(value='Ready. Default OCR interval: 0.75 seconds.')
        ttk.Label(root, textvariable=self.status, wraplength=710, padding=8).pack(fill='x')
        self.preview = ScrolledText(root, wrap='word', font=('Segoe UI', 10), state='disabled')
        self.preview.pack(fill='both', expand=True, padx=12, pady=8)
        root.protocol('WM_DELETE_WINDOW', self.close)
        root.after(100, self.poll)
        root.after(400, self.start)  # one selection at startup

    def open_folder(self):
        if self.output:
            os.startfile(str(self.output.base.parent))

    def start(self):
        from tkinter import filedialog, messagebox
        if self.thread and self.thread.is_alive():
            return
        folder = self.args.output or filedialog.askdirectory(parent=self.root, title='Choose transcript output folder')
        if not folder:
            return
        try:
            output = Output(Path(folder))
            output.save(Transcript())  # verify write access before capture
            region = select_region(self.root)
        except Exception as exc:
            messagebox.showerror('Cannot start', str(exc), parent=self.root)
            return
        if not region:
            self.status.set('Selection cancelled. Click Select region & start to try again.')
            return
        self.output = output
        self.start_button.config(state='disabled')
        self.stop_button.config(state='normal')
        self.folder_button.config(state='normal')
        self.stop_event.clear()
        self.status.set(f'Capturing fixed region {region}. Move this window outside it. Saving to {output.base.parent}')
        self.thread = threading.Thread(target=self.worker, args=(region, output), daemon=False)
        self.root.iconify()  # keep the control window out of the OCR region
        self.thread.start()

    def stop(self):
        self.stop_event.set()
        self.stop_button.config(state='disabled')
        self.status.set('Stopping after the current OCR pass and exporting…')

    def close(self):
        if self.thread and self.thread.is_alive():
            self.closing = True
            self.stop()
        else:
            self.root.destroy()

    def worker(self, region, output):
        import mss
        from PIL import Image
        parser = CaptionParser(self.args.speaker)
        transcript = Transcript(self.args.gap)
        last_docx = time.monotonic()
        failures = 0
        self.stop_event.wait(.4)  # allow the control window to minimize
        try:
            with mss.mss() as capture:
                while not self.stop_event.is_set():
                    began = time.monotonic()
                    observed = datetime.now().astimezone()
                    shot = capture.grab(region)  # fixed coordinates; never reads cursor
                    frame = Image.frombytes('RGB', shot.size, shot.rgb)
                    try:
                        rows = ocr_rows(frame, self.executable, self.args.language, self.args.trim_left)
                    except (RuntimeError, subprocess.TimeoutExpired) as exc:
                        failures += 1
                        self.events.put(('status', f'OCR error ({failures}/3): {exc}'))
                        if failures >= 3:
                            raise RuntimeError('Three consecutive OCR failures; check Tesseract/language installation.') from exc
                        self.stop_event.wait(self.args.interval)
                        continue
                    failures = 0
                    cards = parser.parse(rows)
                    if transcript.consume(cards, observed):
                        output.save(transcript)
                        self.events.put(('text', transcript.render()))
                    if time.monotonic() - last_docx >= 30:
                        output.docx(transcript)
                        last_docx = time.monotonic()
                    elapsed = time.monotonic() - began
                    status = f'{len(transcript.blocks)} spoken blocks | OCR {elapsed:.2f}s | {len(cards)} visible cards'
                    if not cards:
                        status += ' | No named captions: check region and speaker font/--speaker.'
                    if elapsed > self.args.interval:
                        status += ' | OCR slower than requested interval; crop more tightly.'
                    if output.docx_warning:
                        status += ' | ' + output.docx_warning
                    self.events.put(('status', status))
                    self.stop_event.wait(max(0, self.args.interval-elapsed))
        except Exception as exc:
            self.events.put(('error', str(exc)))
        finally:
            try:
                output.save(transcript)
                output.docx(transcript)
                self.events.put(('status', f'Saved: {output.base.with_suffix(".txt")}\n' +
                                 (output.docx_warning or f'Saved: {output.base.with_suffix(".docx")}')))
            except Exception as exc:
                self.events.put(('error', f'Final save failed: {exc}'))
            self.events.put(('done', None))

    def poll(self):
        from tkinter import messagebox
        try:
            while True:
                kind, value = self.events.get_nowait()
                if kind == 'text':
                    self.preview.config(state='normal')
                    self.preview.delete('1.0', 'end')
                    self.preview.insert('end', value)
                    self.preview.see('end')
                    self.preview.config(state='disabled')
                elif kind == 'status':
                    self.status.set(value)
                elif kind == 'error':
                    messagebox.showerror('Capture error', value, parent=self.root)
                elif kind == 'done':
                    self.start_button.config(state='normal')
                    self.stop_button.config(state='disabled')
                    if self.closing:
                        self.root.destroy()
                        return
        except queue.Empty:
            pass
        self.root.after(100, self.poll)


def recover(path: Path):
    latest = {}
    for line in path.read_text(encoding='utf-8').splitlines():
        try:
            state = json.loads(line)
            latest[int(state['block'])] = state
        except (ValueError, KeyError, TypeError):
            continue  # ignore an incomplete last journal record after a crash
    transcript = Transcript()
    for idx in sorted(latest):
        state = latest[idx]
        transcript.blocks.append(Block(state['speaker'], state['start'], state['end'], [Segment(idx, state['text'])]))
    output = Output(path.parent)
    output.base = path.with_name(path.stem + '_recovered')
    atomic_text(output.base.with_suffix('.txt'), transcript.render())
    output.docx(transcript)
    print('Recovered:', output.base.with_suffix('.txt'))
    print(output.docx_warning or f'Recovered: {output.base.with_suffix(".docx")}')


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--interval', type=float, default=.75, help='Requested interval, 0.5–1.0 seconds (default .75)')
    ap.add_argument('--output', type=Path, help='Output folder; otherwise ask at startup')
    ap.add_argument('--tesseract', help='Path to tesseract.exe')
    ap.add_argument('--language', default='eng', help='Installed Tesseract language (default eng)')
    ap.add_argument('--speaker', action='append', default=[], help='Known display name; repeat for multiple speakers')
    ap.add_argument('--trim-left', type=int, default=0, help='Pixels to remove from left of selected region')
    ap.add_argument('--gap', type=float, default=8, help='Seconds without new text before a new block (default 8)')
    ap.add_argument('--recover', type=Path, help='Recover TXT/DOCX from a session JSONL journal')
    args = ap.parse_args()
    if args.recover:
        recover(args.recover)
        return
    if not .5 <= args.interval <= 1 or args.trim_left < 0 or args.gap <= 0:
        ap.error('Use --interval 0.5 to 1.0, --trim-left >= 0 and --gap > 0.')
    if sys.platform != 'win32':
        ap.error('Screen capture UI is intended for Windows. Parser tests/recovery can run elsewhere.')
    dpi_awareness()
    import tkinter as tk
    from tkinter import messagebox
    root = tk.Tk()
    try:
        executable = find_tesseract(args.tesseract)
        import mss
        from PIL import ImageTk
        App(root, args, executable)
    except Exception as exc:
        messagebox.showerror('Setup required', str(exc), parent=root)
        root.destroy()
        return
    root.mainloop()


if __name__ == '__main__':
    main()
